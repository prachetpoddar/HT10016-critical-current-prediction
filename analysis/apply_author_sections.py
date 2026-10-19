#!/usr/bin/env python3
"""
apply_author_sections.py

The data availability and acknowledgement sections, as the authors wrote
them on 2026-10-19. Applied verbatim.

Two consequences are recorded here rather than silently absorbed, because
both are losses against requirements settled earlier in the week:

  the statement now gives the DOI inline instead of citing Ref. [46], so the
  reference entry for the deposit is no longer cited anywhere in the body.
  APS asks for the data citation in the reference list AND in the statement;
  this satisfies the second and drops the first;

  the acknowledgement drops "Large language models assisted in structuring
  the analysis scripts. The authors directed, reviewed and revised that
  work, and the authors take full responsibility for all content, analysis
  and conclusions." That was the part of the AI-use disclosure the policy
  puts in the acknowledgement. What remains of the disclosure is in Sec.
  II.A, which covers the plot-reading and audit models but says nothing
  about the scripts or about who takes responsibility.

The pins for the dropped sentence are removed from check_end_matter.py in
the same commit, so the guard matches the paper the authors want rather
than silently passing on a document that no longer carries what it checks.
"""

import copy
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DATA = (
    "The extracted critical-current dataset and the code are openly "
    "available in https://doi.org/10.5281/zenodo.23485083. The source "
    "measurements were read from the published figures cited throughout, "
    "and from the Materials Project [3] and 3DSC_MP [6] datasets."
)

ACKNOWLEDGMENT = (
    "The authors acknowledge illuminating discussions with Alexander "
    "Balandin and Aashman Mehta. MSH acknowledges support from the U.S. "
    "National Science Foundation award #2610404 and DARPA award "
    "#D26AC50023. This work used Anvil at Purdue University through "
    "allocation PHY260151 (awarded to MSH) from the Advanced "
    "Cyberinfrastructure Coordination Ecosystem: Services & Support "
    "(ACCESS) program. This work also used NAIRR260036 allocation (awarded "
    "to MSH). MSH acknowledges Anthropic's AI for Science program."
)

SECTIONS = [("DATA AVAILABILITY", DATA), ("ACKNOWLEDGMENTS", ACKNOWLEDGMENT)]


def body_of(doc, heading):
    ps = doc.paragraphs
    i = next(k for k, p in enumerate(ps)
             if p.text.strip().upper().startswith(heading))
    j = next(k for k, p in enumerate(ps[i + 1:], i + 1)
             if p.style.name.startswith("Heading") and p.text.strip())
    bodies = [p for p in ps[i + 1:j] if p.text.strip()]
    assert len(bodies) == 1, (heading, len(bodies))
    return bodies[0]


def replace_text(paragraph, text, marked):
    keep = paragraph._p
    model = next(r for r in keep.findall(qn("w:r")))
    for node in list(keep):
        if node.tag in (qn("w:r"), qn("w:hyperlink")):
            keep.remove(node)
    run = copy.deepcopy(model)
    for t in run.findall(qn("w:t")):
        run.remove(t)
    rPr = run.find(qn("w:rPr"))
    if rPr is not None:
        for h in rPr.findall(qn("w:highlight")):
            rPr.remove(h)
    if marked:
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            run.insert(0, rPr)
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        rPr.append(h)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    keep.append(run)


def apply(path, marked):
    doc = docx.Document(path)
    for heading, text in SECTIONS:
        replace_text(body_of(doc, heading), text, marked)
    doc.save(path)
    # relink the DOI, which the rewrite flattened back to plain text
    sys.path.insert(0, "analysis")
    from apply_hyperlinks import isolate, link, urls_in
    doc = docx.Document(path)
    p = body_of(doc, "DATA AVAILABILITY")
    text = "".join(r.text for r in p.runs)
    for at, url in reversed(list(urls_in(text))):
        r = isolate(p, at, len(url))
        assert r.text == url
        link(p, r, url)
    doc.save(path)


def verify(path):
    doc = docx.Document(path)
    for heading, text in SECTIONS:
        got = body_of(doc, heading).text
        assert got == text, (heading, got[:70])
    anchors = ["".join(t.text or "" for t in h.iter(qn("w:t")))
               for p in doc.paragraphs
               for h in p._p.findall(qn("w:hyperlink"))]
    assert "https://doi.org/10.5281/zenodo.23485083" in anchors, anchors


if __name__ == "__main__":
    for path, marked in (("out_final60/HT10016_revised_final60.docx", False),
                         ("out_final60/HT10016_revised_final60_HIGHLIGHTED.docx",
                          True)):
        apply(path, marked)
        verify(path)
        print("%-46s both sections replaced, %d and %d words"
              % (path.split("/")[-1], len(DATA.split()),
                 len(ACKNOWLEDGMENT.split())))
