#!/usr/bin/env python3
"""
apply_ai_disclosure.py

Shorten the artificial-intelligence use disclosure.

It ran to four paragraphs and 306 words. APS prescribes no length. What it
prescribes is a content floor, the tool name and version, how the AI
assisted, and how the authors directed and verified its output, and a
placement rule that puts AI used in conducting research where the methods are
described. Section II.A already names the two plot readers, so the standalone
statement only has to carry the floor.

Four phrases are kept word for word rather than rewritten, because each one
settles something:

  "Curve-level extraction was performed by a single model, GPT-4o (OpenAI)"
  "A second independent reader, Claude Opus 4.8 (Anthropic)"
      the naming T. Chowdhury asked for, which Sec. II.A now matches;

  "the first author checked the outputs by hand ... for all 23 fully fittable
  compounds of the pre-withdrawal cohort and for every disagreement the
  second reader raised"
      the answer to her question about who checked, and its scope. This one
      does not stay in the disclosure. It moves into Sec. II.A, beside the
      sentence that already escalates reader disagreements to manual reading,
      which is where APS puts AI use in conducting research and where a
      reader looking for the verification protocol will be;

  "The authors directed, reviewed and revised that work" and "the authors
  take full responsibility for all content, analysis and conclusions"
      the accountability APS asks for, and deliberately plural.

So this lands at four sentences and about 115 words, against 306 before.
"""

import copy
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, "analysis")
from apply_data_availability import section_bounds as _bounds  # noqa: E402

HEADING = "ARTIFICIAL INTELLIGENCE USE DISCLOSURE"
NEXT = "DATA AVAILABILITY"

NEW = (
    "Curve-level extraction was performed by a single model, GPT-4o "
    "(OpenAI) [27, 28]. A second independent reader, Claude Opus 4.8 "
    "(Anthropic) [36], read the fitted cohort and the critical-field layer, "
    "and the verification audit of Section II.A was run by Claude Opus 5 and "
    "Claude Sonnet 5 (Anthropic) [36], under fixed per-figure prompts that "
    "are deposited with the data. The models located axes and scales, "
    "digitized data-point coordinates, and read printed labels and captions; "
    "they performed no fitting, aggregation, validation or prediction, and "
    "Section II.A states how their output was verified. Large language "
    "models assisted in structuring the analysis scripts, the authors "
    "directed, reviewed and revised that work, and the authors take full "
    "responsibility for all content, analysis and conclusions."
)

# The sentence that leaves the disclosure, and where it lands. The anchor is
# the sentence in Sec. II.A that already escalates a disagreement between the
# two readers to manual reading, so the named checker and the scope of the
# check sit with the protocol they belong to.
IIA_ANCHOR = ("Disagreements between the two readers were escalated to "
              "manual reading against the source figure.")
IIA_ADDED = (" The first author checked the outputs by hand against the "
             "source figures for all 23 fully fittable compounds of the "
             "pre-withdrawal cohort and for every disagreement the second "
             "reader raised.")

KEPT = [
    "Curve-level extraction was performed by a single model, GPT-4o (OpenAI)",
    "A second independent reader, Claude Opus 4.8 (Anthropic)",
    "the authors directed, reviewed and revised that work",
    "the authors take full responsibility for all content, analysis and "
    "conclusions",
]

# These two have to survive the move, somewhere in the manuscript.
MOVED = ["the first author checked the outputs by hand against the source "
         "figures",
         "for all 23 fully fittable compounds of the pre-withdrawal cohort "
         "and for every disagreement the second reader raised"]


def bounds(doc):
    ps = doc.paragraphs
    start = next(i for i, p in enumerate(ps)
                 if p.text.strip().upper().startswith(HEADING))
    end = next(i for i, p in enumerate(ps[start + 1:], start + 1)
               if p.text.strip().upper().startswith(NEXT))
    return start, end


def move_the_check(path):
    """Put the manual-check sentence into Sec. II.A, once."""
    from runs import replace
    doc = docx.Document(path)
    if any(IIA_ADDED.strip() in p.text for p in doc.paragraphs):
        return False
    hit = [p for p in doc.paragraphs if IIA_ANCHOR in p.text]
    assert len(hit) == 1, (path, len(hit))
    replace(hit[0], IIA_ANCHOR, IIA_ANCHOR + IIA_ADDED)
    doc.save(path)
    return True


def apply(path, marked):
    doc = docx.Document(path)
    start, end = bounds(doc)
    body = doc.paragraphs[start + 1:end]
    first = next(p for p in body if p.text.strip())
    keep = first._p
    model = keep.findall(qn("w:r"))[0]
    for node in list(keep):
        if node.tag in (qn("w:r"), qn("w:hyperlink")):
            keep.remove(node)
    run = copy.deepcopy(model)
    for t in run.findall(qn("w:t")):
        run.remove(t)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = NEW
    run.append(t)
    if marked:
        rPr = run.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            run.insert(0, rPr)
        for old in rPr.findall(qn("w:highlight")):
            rPr.remove(old)
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        rPr.append(h)
    keep.append(run)
    removed = 0
    for p in body:
        if p._p is keep:
            continue
        p._p.getparent().remove(p._p)
        removed += 1
    doc.save(path)
    return removed


def verify(path, marked):
    doc = docx.Document(path)
    start, end = bounds(doc)
    body = [p for p in doc.paragraphs[start + 1:end] if p.text.strip()]
    assert len(body) == 1, "%d paragraph(s)" % len(body)
    assert body[0].text == NEW, body[0].text[:70]
    for phrase in KEPT:
        assert phrase in NEW, phrase[:40]
    # the sentence is capitalised where it now begins, so compare case-blind
    whole = " ".join(p.text for p in doc.paragraphs).lower()
    for phrase in MOVED:
        assert phrase.lower() in whole, "lost in the move: %s" % phrase[:40]
        assert phrase.lower() not in NEW.lower(), phrase[:40]
    if marked:
        from docx.enum.text import WD_COLOR_INDEX
        assert all(r.font.highlight_color == WD_COLOR_INDEX.YELLOW
                   for r in body[0].runs if r.text.strip())


if __name__ == "__main__":
    import re
    for path, marked in (("out_final60/HT10016_revised_final60.docx", False),
                         ("out_final60/HT10016_revised_final60_HIGHLIGHTED.docx",
                          True)):
        moved = move_the_check(path)
        removed = apply(path, marked)
        verify(path, marked)
        print("%-46s manual-check sentence %s Sec. II.A"
              % (path.split("/")[-1],
                 "moved into" if moved else "already in"))
        print("%-46s %d paragraph(s) removed, now %d words in %d sentences"
              % (path.split("/")[-1], removed, len(NEW.split()),
                 len(re.findall(r"[.!?](?:\s|$)", NEW))))
