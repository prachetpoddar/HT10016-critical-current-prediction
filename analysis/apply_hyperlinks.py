#!/usr/bin/env python3
"""
apply_hyperlinks.py

Turn every bare URL in the delivered documents into a real Word hyperlink, so a
reader can click it rather than copy it out of the text.

Requested by T. Chowdhury on the final59 draft: "Please consider making any
URLs (e.g., GitHub links) clickable in the manuscript and other files."

The visual appearance of the URL is left exactly as it was. Only the
relationship is added, so the text stays Times New Roman 12 in the body and
keeps its own run properties in the reference list. Word renders the span as a
link because of the w:hyperlink wrapper, not because of a colour.

A trailing sentence period is not part of the URL and is left outside the link.
"""

import copy
import re
import sys

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

URL = re.compile(r"https?://[^\s<>()\[\]]+")
TRAILING = ".,;:"


def urls_in(text):
    for m in URL.finditer(text):
        url = m.group(0)
        while url and url[-1] in TRAILING:
            url = url[:-1]
        if url:
            yield m.start(), url


def isolate(paragraph, at, length):
    """Split runs so that the span [at, at+length) is exactly one run, and
    return it. Every split run keeps a copy of the original run properties."""
    end = at + length
    pos = 0
    for run in list(paragraph.runs):
        start, stop = pos, pos + len(run.text)
        pos = stop
        if stop <= at or start >= end:
            continue
        head = run.text[: max(0, at - start)]
        body = run.text[max(0, at - start) : min(len(run.text), end - start)]
        tail = run.text[min(len(run.text), end - start) :]
        if head:
            before = copy.deepcopy(run._r)
            run._r.addprevious(before)
            _settext(before, head)
        run.text = body
        if tail:
            after = copy.deepcopy(run._r)
            run._r.addnext(after)
            _settext(after, tail)
        return run
    raise ValueError("no run holds the span at %d" % at)


def _settext(r, text):
    for t in r.findall(qn("w:t")):
        r.remove(t)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)


def link(paragraph, run, url):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run._r.addprevious(h)
    h.append(run._r)
    return h


def apply(path):
    doc = docx.Document(path)
    done = []
    for paragraph in doc.paragraphs:
        text = "".join(r.text for r in paragraph.runs)
        if "http" not in text:
            continue
        # Work right to left so earlier offsets stay valid after a split.
        for at, url in reversed(list(urls_in(text))):
            run = isolate(paragraph, at, len(url))
            assert run.text == url, (run.text, url)
            link(paragraph, run, url)
            done.append(url)
    doc.save(path)
    return done


def verify(path, expected):
    doc = docx.Document(path)
    rels = [
        r for r in doc.part.rels.values()
        if r.reltype == RT.HYPERLINK and r.is_external
    ]
    anchors = []
    for p in doc.paragraphs:
        for h in p._p.findall(qn("w:hyperlink")):
            anchors.append("".join(t.text or "" for t in h.iter(qn("w:t"))))
    assert sorted(anchors) == sorted(expected), (anchors, expected)
    for a in anchors:
        assert a in [r.target_ref for r in rels], a
    return anchors


if __name__ == "__main__":
    for path in sys.argv[1:]:
        before = docx.Document(path)
        text_before = [p.text for p in before.paragraphs]
        done = apply(path)
        after = docx.Document(path)
        assert [p.text for p in after.paragraphs] == text_before, path
        verify(path, done)
        print("%-46s %d link(s)" % (path.split("/")[-1], len(done)))
        for u in done:
            print("    " + u)
