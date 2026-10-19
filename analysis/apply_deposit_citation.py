#!/usr/bin/env python3
"""
apply_deposit_citation.py

Cite the deposit as a reference, and shorten the data statement to match.

APS requires that publicly shared data be cited in the reference list and
that the same citation appear in the data availability statement, with
creators, year, version, repository and a persistent identifier. The
statement carried a bare GitHub URL and the reference list carried no entry
for the deposit at all, so neither half of that was satisfied. The Zenodo
release of 2026-10-19 supplies the identifier.

The publisher list goes. Elsevier, Springer Nature, IOP Publishing and arXiv
were 30 words restating what references [9] to [44] already say, and APS does
not ask for third-party publishers of articles the authors read data out of.
What it does ask for is a citation of reused third-party data, which is why
the Materials Project and 3DSC_MP keep their references here.

Not closed by this pass: the nine field-axis records digitized from National
High Magnetic Field Laboratory datasets are cited nowhere. The deposit
identifies them only by internal names such as MAGLAB_Ni122_4_2K, with no URL
or dataset title to cite, so the source has to come from the authors.

The Zenodo record itself could not be read when this was written, because the
fetch was rate-limited. The creator list, title, version and year below are
the ones the release was configured with and should be checked against the
record before submission.
"""

import copy
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, "analysis")
from apply_hyperlinks import isolate, link, urls_in  # noqa: E402

DOI = "https://doi.org/10.5281/zenodo.23485083"

REFERENCE = (
    "[46] P. Poddar, T. Chowdhury, and M. S. Hossain, "
    "Substructure-conditional critical current density prediction: data and "
    "analysis code, Version v1.0.0, Zenodo (2026), " + DOI + "."
)

NEW_DAS = (
    "The extracted critical-current dataset and the code that regenerates "
    "every number in this paper are openly available in Ref. [46]. The "
    "source measurements were read from the published figures cited "
    "throughout, and from the Materials Project [3] and 3DSC_MP [6] "
    "datasets."
)


def das_bounds(doc):
    ps = doc.paragraphs
    start = next(i for i, p in enumerate(ps)
                 if p.text.strip().upper().startswith("DATA AVAILABILITY"))
    end = next(i for i, p in enumerate(ps[start + 1:], start + 1)
               if p.text.strip().upper().startswith("REFERENCES"))
    return start, end


def rewrite_das(doc, marked):
    start, end = das_bounds(doc)
    paragraph = next(p for p in doc.paragraphs[start + 1:end] if p.text.strip())
    keep = paragraph._p
    model = None
    for node in list(keep):
        if node.tag == qn("w:r") and model is None:
            model = node
        if node.tag in (qn("w:r"), qn("w:hyperlink")):
            keep.remove(node)
    run = copy.deepcopy(model)
    for t in run.findall(qn("w:t")):
        run.remove(t)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = NEW_DAS
    run.append(t)
    if marked:
        rPr = run.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            run.insert(0, rPr)
        for old in rPr.findall(qn("w:highlight")):
            rPr.remove(old)
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        rPr.append(h)
    keep.append(run)


def add_reference(doc):
    """Append the entry after [45], copying that paragraph's formatting."""
    last = [p for p in doc.paragraphs if p.text.strip().startswith("[45]")]
    assert len(last) == 1, len(last)
    if any(p.text.strip().startswith("[46]") for p in doc.paragraphs):
        return False
    new = copy.deepcopy(last[0]._p)
    for node in list(new):
        if node.tag in (qn("w:r"), qn("w:hyperlink")):
            new.remove(node)
    run = copy.deepcopy(last[0]._p.findall(qn("w:r"))[0])
    for t in run.findall(qn("w:t")):
        run.remove(t)
    for rPr in run.findall(qn("w:rPr")):
        for h in rPr.findall(qn("w:highlight")):
            rPr.remove(h)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = REFERENCE
    run.append(t)
    new.append(run)
    last[0]._p.addnext(new)
    return True


def relink(doc):
    for paragraph in doc.paragraphs:
        text = "".join(t.text or "" for t in paragraph._p.iter(qn("w:t")))
        if DOI not in text or paragraph._p.findall(qn("w:hyperlink")):
            continue
        for at, url in reversed(list(urls_in(text))):
            r = isolate(paragraph, at, len(url))
            assert r.text == url, (r.text, url)
            link(paragraph, r, url)


def verify(path):
    doc = docx.Document(path)
    ps = [p.text.strip() for p in doc.paragraphs]
    assert REFERENCE in ps, "the reference entry is not there"
    assert ps.index(REFERENCE) > ps.index(
        next(t for t in ps if t.startswith("[45]"))), "out of order"
    start, end = das_bounds(doc)
    body = [p for p in doc.paragraphs[start + 1:end] if p.text.strip()]
    assert len(body) == 1 and body[0].text == NEW_DAS, body[0].text[:60]
    for gone in ("sciencedirect", "link.springer", "iopscience", "arxiv.org",
                 "nationalmaglab", "materialsproject.org"):
        assert gone not in body[0].text, gone
    anchors = [t.text or "" for p in doc.paragraphs
               for h in p._p.findall(qn("w:hyperlink"))
               for t in h.iter(qn("w:t"))]
    assert DOI in anchors, anchors
    rels = [r.target_ref for r in doc.part.rels.values()
            if r.reltype.endswith("/hyperlink")]
    assert DOI in rels


if __name__ == "__main__":
    for path, marked in (("out_final60/HT10016_revised_final60.docx", False),
                         ("out_final60/HT10016_revised_final60_HIGHLIGHTED.docx",
                          True)):
        doc = docx.Document(path)
        added = add_reference(doc)
        rewrite_das(doc, marked)
        doc.save(path)
        doc = docx.Document(path)
        relink(doc)
        doc.save(path)
        verify(path)
        print("%-46s reference [46] %s, statement now %d words"
              % (path.split("/")[-1], "added" if added else "already present",
                 len(NEW_DAS.split())))
