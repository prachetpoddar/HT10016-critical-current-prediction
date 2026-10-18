#!/usr/bin/env python3
"""
clear_highlighting.py

Take the revision marking off a document, so the clean copy is clean.

Rewritten on 2026-10-18 after the working container was reclaimed. The marked
copy is what the editor reads and what the collaborators comment on; the file
the journal receives carries no highlighting. Every colour found is reported,
since a colour other than yellow would mean something other than revision
marking.

    python3 analysis/clear_highlighting.py IN.docx OUT.docx
"""
import shutil
import sys
from collections import Counter

import docx


def stream(doc):
    for para in doc.paragraphs:
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield para


def main(src, dst):
    if src != dst:
        shutil.copyfile(src, dst)
    doc = docx.Document(dst)
    seen, chars = Counter(), 0
    for para in stream(doc):
        for run in para.runs:
            if run.font.highlight_color is None:
                continue
            seen[str(run.font.highlight_color).split(" ")[0]] += 1
            chars += len(run.text)
            run.font.highlight_color = None
    doc.save(dst)

    check = docx.Document(dst)
    left = sum(1 for p in stream(check) for r in p.runs
               if r.font.highlight_color is not None)
    if left:
        sys.exit("%d highlighted run(s) remain in %s" % (left, dst))
    a = " ".join(p.text for p in docx.Document(src).paragraphs).split()
    b = " ".join(p.text for p in check.paragraphs).split()
    if a != b:
        sys.exit("clearing the highlighting changed the text of %s" % dst)
    print("   %-52s cleared %d chars, %s"
          % (dst.split("/")[-1], chars,
             ", ".join("%s x%d" % (k, v) for k, v in seen.items()) or "none"))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
