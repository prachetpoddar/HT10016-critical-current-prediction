#!/usr/bin/env python3
"""
check_review_items.py

Every item the co-authors asked for, checked against the delivered documents.

The list is theirs, not mine: T. Chowdhury's two rounds of comments, and
M. S. Hossain's font, figure, table, spacing and paragraph-length items. Each
is expressed as text that must be present or absent, or as a property of the
file, so a later edit that quietly undoes one of them fails here rather than
reaching a referee.

Pins move when a decision moves. The model namings were pinned to the AI
disclosure until 2026-10-19, when that section was removed and its content
placed in Sec. II.A and the acknowledgement; the pins now point at the
sentences in Sec. II.A that carry the same requirement. A pin that is
repointed without the requirement surviving somewhere is the failure this
file exists to prevent, so the new target is the co-author's requirement
rather than the nearest surviving wording.

    python3 analysis/check_review_items.py
"""

import docx, re, glob, json, os, zipfile
from docx.oxml.ns import qn

OUT = "out_final60"
def running(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts += [p.text for p in c.paragraphs]
    return re.sub(r"\s+", " ", " ".join(parts))

MAN = running(OUT+"/HT10016_revised_final60.docx")
SUP = running(OUT+"/SUPPLEMENTAL_MATERIAL_revised_final60.docx")
LET = running(OUT+"/RESPONSE_TO_REFEREES_final60.docx")
COV = running(OUT+"/COVER_LETTER_final60.docx")
def norm(s): return re.sub(r"\s+"," ",s)

rows=[]
def check(owner, item, ok, detail=""):
    rows.append((owner, item, ok, detail))

# ---------- Chowdhury, first round (apply_coauthor_review.py) ----------
GONE = ["were the normalization scales measured","assigns two distinct regimes",
        "when to stay silent","Our goal is concrete.","Here each record is a curve. It",
        "a different organization and a different generation"]
KEPT = ["whether the normalization scales are measured on the same specimen",
        "It assigns three distinct regimes across the studied families.",
        "The diagnostic yields three distinct outcomes",
        "Above 0.7 the family is treated as sample-form dominant, from 0.3 to 0.7 as moderate, and below 0.3 as minor.",
        "when to refrain from predicting","Our concrete goal is to predict",
        "Here each record is a curve, tied to a source paper",
        "The first model, GPT-4o (OpenAI), read every figure",
        "Claude Opus 5 and Claude Sonnet 5 (Anthropic) [36]",
        "The first model, GPT-4o (OpenAI), read every figure",
        "A second reader checks the reliability of that extraction"]
labels = {0:"1 abstract both questions present tense",1:"6 conclusion regime count three",
          2:"2 refrain from predicting",3:"3 goal in one sentence",4:"4 record in one sentence",
          5:"5 readers named, 'company'"}
for i,g in enumerate(GONE):
    check("Chowdhury 1", "superseded text gone: "+labels[i], norm(g) not in MAN, g[:40])
for k in KEPT:
    check("Chowdhury 1", "new text present: "+k[:52], norm(k) in MAN)

# the manual-check attribution
check("Chowdhury 1","manual check attributed to the first author",
      "first author checked the outputs by hand" in MAN.replace("The first","the first")
      and "the authors checked the outputs by hand" not in MAN)
for k in [
          "reported by the authors",
          "for all 23 fully fittable compounds of the pre-withdrawal cohort"]:
    check("Chowdhury 1","plural use kept: "+k[:46], norm(k) in MAN)

# ---------- Chowdhury, second round ----------
def links(path):
    d=docx.Document(path)
    out=[]
    for p in d.paragraphs:
        for h in p._p.findall(qn("w:hyperlink")):
            out.append("".join(t.text or "" for t in h.iter(qn("w:t"))))
    rels=[r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    return out, rels
URLRE=re.compile(r"https?://[^\s<>()\[\]]+")
for name in sorted(os.path.basename(f) for f in glob.glob(OUT+"/*.docx")):
    text=running(OUT+"/"+name)
    urls={u.rstrip(".,;:") for u in URLRE.findall(text)}
    anch,rels=links(OUT+"/"+name)
    ok = urls == set(anch) and all(a in rels for a in anch)
    check("Chowdhury 2","every URL clickable in "+name[:34], ok,
          "%d url(s), %d link(s)" % (len(urls), len(anch)))
readme=open("README.md").read()
check("Chowdhury 2","README landing section at the top",
      readme.index("## Where everything is") < readme.index("## Reproducing"))
check("Chowdhury 2","README directory structure", "HT10016-critical-current-prediction/" in readme and "data/" in readme)
check("Chowdhury 2","README maps paper sections to files", "## Which part of the paper rests on which files" in readme)

# ---------- Hossain, earlier ----------
for name in sorted(os.path.basename(f) for f in glob.glob(OUT+"/*.docx")):
    d=docx.Document(OUT+"/"+name)
    styles=zipfile.ZipFile(OUT+"/"+name).read("word/styles.xml").decode("utf8")
    head=styles.split("</w:docDefaults>")[0]
    body_runs=[r for p in d.paragraphs for r in p.runs if r.text.strip()]
    bad_font=sorted({r.font.name for r in body_runs if r.font.name and r.font.name not in
                     ("Times New Roman","Cambria Math","Courier New","Consolas","Menlo")})
    bad_size=sorted({r.font.size.pt for r in body_runs if r.font.size and r.font.size.pt!=12})
    check("Hossain earlier","Times New Roman 12 body in "+name[:32],
          'w:sz w:val="24"' in head and not bad_font and not bad_size,
          "stray fonts %s, stray sizes %s" % (bad_font or "none", bad_size or "none"))
check("Hossain earlier","referee acknowledgement removed from the manuscript",
      not re.search(r"we thank the referee|thank the referees|referees for their", MAN, re.I))
check("Hossain earlier","the letter still thanks the referees",
      bool(re.search(r"thank both referees|thank the referees", LET, re.I)))

# ---------- Hossain, the six formatting items ----------
d=docx.Document(OUT+"/HT10016_revised_final60.docx")
shapes=list(d.inline_shapes)
w=[s.width.inches for s in shapes]
check("Hossain format","1 Figure 4 is the smallest drawing", w[3]==min(w) and w[3]<6.0,
      "widths %s in" % ["%.2f"%x for x in w])
src=open("analysis/manuscript_figure_4.py").read()
check("Hossain format","2 Figure 4 axis labels capitalised",
      'set_xlabel("Anchor measurements' in src and 'set_ylabel("Pooled error' in src)
bad=[]
for name in ["HT10016_revised_final60.docx","SUPPLEMENTAL_MATERIAL_revised_final60.docx"]:
    doc=docx.Document(OUT+"/"+name)
    for i,t in enumerate(doc.tables):
        pr=t._tbl.find(qn("w:tblPr")); b=pr.find(qn("w:tblBorders")) if pr is not None else None
        got={e.tag.split('}')[1]: e.get(qn("w:val")) for e in (b if b is not None else [])}
        if set(got)!={"top","left","bottom","right","insideH","insideV"} or set(got.values())!={"single"}:
            bad.append((name[:20],i,got))
check("Hossain format","3 every table has visible borders on all edges", not bad, str(bad)[:60])
sizes={}
for name in ["HT10016_revised_final60.docx","SUPPLEMENTAL_MATERIAL_revised_final60.docx"]:
    doc=docx.Document(OUT+"/"+name)
    s=sorted({r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells
              for p in c.paragraphs for r in p.runs if r.font.size})
    sizes["man" if name.startswith("HT") else "sup"]=s
check("Hossain format","4 table text one point down",
      sizes["man"]==[11.0] and sizes["sup"]==[9.5,11.0], str(sizes))
plan=json.load(open("/tmp/m_final.json")) if os.path.exists("/tmp/m_final.json") else None
check("Hossain format","5 ceiling: no body paragraph over ten lines", plan==[] if plan is not None else None,
      "see planner run")
for name in ["HT10016_revised_final60.docx","HT10016_revised_final60_HIGHLIGHTED.docx",
             "SUPPLEMENTAL_MATERIAL_revised_final60.docx","SUPPLEMENTAL_MATERIAL_revised_final60_HIGHLIGHTED.docx"]:
    doc=docx.Document(OUT+"/"+name)
    body=doc.element.body
    P=qn("w:p"); TBL=qn("w:tbl")
    def ntext(n): return "".join(t.text or "" for t in n.iter(qn("w:t"))).strip()
    seq=[]
    for n in body:
        if n.tag==TBL: seq.append(("table",""))
        elif n.tag==P: seq.append(("text" if ntext(n) else "blank", ntext(n)))
    start=next(i for i,(k,t) in enumerate(seq) if t.startswith("ABSTRACT") or t.startswith("1. "))
    stop=next((i for i,(k,t) in enumerate(seq) if t.upper().startswith("REFERENCES")), len(seq))
    runs=0; dbl=0; missing=0; prev=None
    for kind,t in seq[start:stop]:
        if kind=="table":
            prev="table"; runs=0; continue
        if kind=="text":
            if prev=="text": missing+=1
            prev="text"; runs=0
        else:
            runs+=1; dbl += runs>1; prev="blank"
    check("Hossain format","6 exactly one blank line in "+name[:30], missing==0 and dbl==0,
          "%d missing, %d doubled" % (missing,dbl))

print("%-16s %-62s %s" % ("WHOSE","ITEM","RESULT"))
print("-"*104)
fails=0
for owner,item,ok,detail in rows:
    mark = "ok" if ok else ("?" if ok is None else "FAIL")
    if ok is False: fails+=1
    print("%-16s %-62s %-5s %s" % (owner,item[:62],mark,detail[:28]))
print("\n%d checks, %d failing" % (len(rows), fails))
