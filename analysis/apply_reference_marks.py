#!/usr/bin/env python3
"""
apply_reference_marks.py

A reference entry is marked only when the entry is new. Otherwise it carries
no mark at all.

The marked copy had 88 highlighted fragments spread over 37 of the 46 entries,
and they were not editorial. Their shapes give them away: the bracketed number
plus the first author's initial, the volume, page and year, bare numbers, the
odd journal abbreviation. No author name, title or journal title was marked
anywhere. That is a token-level diff seeing a renumbered list, marking the
numbers that moved and leaving the matching words alone, rather than anyone
recording a change to a reference.

Which entries are new is decided by comparing the reference list against
out_final52, the submitted set in the deposit, on the text of each entry with
its bracket number stripped and case and punctuation flattened, so that a
renumbering alone never reads as a change. Exactly one entry is new, [46], the
deposit itself. No entry in the submitted set is missing from the current one.

The baseline is the submitted set, which is the only earlier full set this
repository holds. If a reference was added between the first submission and
that set, it reads as old here.
"""

import re
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASELINE = "out_final52/HT10016_revised_final52.docx"
TARGETS = ["out_final60/HT10016_revised_final60_HIGHLIGHTED.docx"]
YELLOW = "yellow"


def entries(path):
    doc = docx.Document(path)
    ps = [p.text.strip() for p in doc.paragraphs]
    i = next(k for k, t in enumerate(ps) if t.upper().startswith("REFERENCES"))
    out = {}
    for t in ps[i:]:
        m = re.match(r"^\[(\d+)\]\s*(.+)$", t)
        if m:
            out[int(m.group(1))] = re.sub(r"\s+", " ", m.group(2)).strip()
    return out


def key(s):
    return " ".join(re.sub(r"[^a-z0-9 ]", " ", s.lower()).split())[:90]


def new_entries():
    old = {key(v) for v in entries(BASELINE).values()}
    return {n for n, t in entries(TARGETS[0]).items() if key(t) not in old}


def set_mark(run_el, on):
    rPr = run_el.find(qn("w:rPr"))
    if rPr is None:
        if not on:
            return False
        rPr = OxmlElement("w:rPr")
        run_el.insert(0, rPr)
    had = rPr.find(qn("w:highlight")) is not None
    for old in rPr.findall(qn("w:highlight")):
        rPr.remove(old)
    if on:
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), YELLOW)
        rPr.append(h)
        return not had
    return had


def body_yellow(doc, upto):
    n = 0
    for p in doc.paragraphs[:upto]:
        for r in p._p.iter(qn("w:r")):
            rPr = r.find(qn("w:rPr"))
            if rPr is None:
                continue
            h = rPr.find(qn("w:highlight"))
            if h is not None and h.get(qn("w:val")) == YELLOW:
                n += sum(len(t.text or "") for t in r.iter(qn("w:t")))
    return n


def apply(path, new):
    doc = docx.Document(path)
    ps = doc.paragraphs
    refs_at = next(i for i, p in enumerate(ps)
                   if p.text.strip().upper().startswith("REFERENCES"))
    before_body = body_yellow(doc, refs_at)
    before_text = [p.text for p in ps]
    marked = cleared = 0
    for p in ps[refs_at:]:
        m = re.match(r"^\[(\d+)\]", p.text.strip())
        if not m:
            continue
        on = int(m.group(1)) in new
        for r in p._p.iter(qn("w:r")):
            if not "".join(t.text or "" for t in r.iter(qn("w:t"))):
                continue
            changed = set_mark(r, on)
            if changed and on:
                marked += 1
            elif changed:
                cleared += 1
    doc.save(path)

    doc = docx.Document(path)
    assert [p.text for p in doc.paragraphs] == before_text, "text changed"
    assert body_yellow(doc, refs_at) == before_body, "the body's marks moved"
    for p in doc.paragraphs[refs_at:]:
        m = re.match(r"^\[(\d+)\]", p.text.strip())
        if not m:
            continue
        want = int(m.group(1)) in new
        runs = [r for r in p._p.iter(qn("w:r"))
                if "".join(t.text or "" for t in r.iter(qn("w:t")))]
        for r in runs:
            rPr = r.find(qn("w:rPr"))
            h = None if rPr is None else rPr.find(qn("w:highlight"))
            got = h is not None and h.get(qn("w:val")) == YELLOW
            assert got == want, (p.text[:16], got, want)
    return marked, cleared, before_body


if __name__ == "__main__":
    new = new_entries()
    print("new reference entries: %s" % (sorted(new) or "none"))
    for path in TARGETS:
        marked, cleared, body = apply(path, new)
        print("%-46s %d run(s) marked, %d cleared, body unchanged at %d chars"
              % (path.split("/")[-1], marked, cleared, body))
