#!/usr/bin/env python3
"""
embed_current_figures.py

Put the deposited figure PNGs into the manuscript.

Figures 1 and 2 in the document were an older rendering than the one in
figures/. The deposit's Figure 1 prints "96 per-paper anchors over 32
papers"; the one embedded in the document printed 29, which is not the paper
count of any cohort in the deposit. Every version of the 96-row anchor table
in data/ and in audit/ gives 32 distinct paper_id values, and the generator
reads that count rather than carrying a literal. The embedded Figure 2 was an
earlier design of the same diagram, at a different size, and printed the same
29.

So the figures were regenerated and committed and never re-embedded. This
pass closes that gap for all five, keeping each figure's printed WIDTH and
recomputing its height from the committed image, because Figure 2's new
rendering has a slightly different aspect ratio and Figure 4 is deliberately
drawn narrower than the rest.
"""

import io
import sys

import docx
from docx.shared import Emu
from PIL import Image

TARGETS = ("out_final60/HT10016_revised_final60.docx",
           "out_final60/HT10016_revised_final60_HIGHLIGHTED.docx")

# Document order. check_figures.py uses the same mapping.
FIGURES = ["figures/manuscript_figure_1.png",
           "figures/manuscript_figure_2.png",
           "figures/manuscript_figure_3.png",
           "figures/figure_4_anchor_count.png",
           "figures/manuscript_figure_5.png"]


def aspect(blob):
    with Image.open(io.BytesIO(blob)) as im:
        return im.size[0] / im.size[1], im.size


def apply(path):
    doc = docx.Document(path)
    shapes = list(doc.inline_shapes)
    assert len(shapes) == len(FIGURES), len(shapes)
    changed = []
    for shape, source in zip(shapes, FIGURES):
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        part = doc.part.related_parts[rid]
        want = open(source, "rb").read()
        ratio, size = aspect(want)
        width = shape.width
        height = Emu(int(round(width / ratio)))
        if part.blob == want and abs(shape.height - height) <= 1:
            continue
        part._blob = want
        shape.width, shape.height = width, height
        changed.append((source, size, width.inches, height.inches))
    doc.save(path)
    return changed


def verify(path):
    doc = docx.Document(path)
    for shape, source in zip(doc.inline_shapes, FIGURES):
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        blob = doc.part.related_parts[rid].blob
        assert blob == open(source, "rb").read(), source
        ratio, _ = aspect(blob)
        assert abs(shape.width / ratio - shape.height) <= 1, source


if __name__ == "__main__":
    for path in TARGETS:
        text = [p.text for p in docx.Document(path).paragraphs]
        changed = apply(path)
        verify(path)
        assert [p.text for p in docx.Document(path).paragraphs] == text
        print("%-46s %d figure(s) re-embedded" % (path.split("/")[-1],
                                                  len(changed)))
        for source, size, w, h in changed:
            print("    %-40s %sx%s at %.2f x %.2f in"
                  % (source.split("/")[-1], size[0], size[1], w, h))
