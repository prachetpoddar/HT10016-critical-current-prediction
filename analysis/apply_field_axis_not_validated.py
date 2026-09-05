#!/usr/bin/env python3
"""Report the field axis as not validated at family level.

The decision this implements: stop assigning per-family verdicts to the field
axis. Three findings put it beyond a renumbering.

  The cohort that survives the anchor repairs and the stated fitting protocol
  is 52 fits from 12 papers, against 94 from 16 as published.

  The fitted field exponent does not separate substructure families on it. The
  fraction of between-paper variance the family label accounts for is 0.038
  with a permutation p of 0.98, and 0.055 with p 0.93 on the published cohort,
  so this is not a consequence of the repair either.

  The per-family leave-one-compound-out errors reverse. MgB2-class moves from
  0.753 to 1.230 and iron pnictide 122-type from 0.973 to 1.957, while iron
  chalcogenide 11-type moves from 1.093 to 0.707. The two families that cleared
  the screening-grade threshold on the published anchors stop clearing and the
  one that failed starts.

Source is out_blockers/*_final2.docx. Output goes to out_field/.

What this deliberately does NOT do. Section III.E restricts dispatch to three
families "that pass the validation and population requirements", and the only
family that dispatches anything is MgB2-class, whose validating axis was the
field axis. Removing the field-axis verdicts therefore leaves the dispatch
without a validating axis, and repairing that is a decision about what the
paper claims rather than a correction to what it reports. The passages are
located and printed on every run so that the inconsistency is visible rather
than silently carried, and the script refuses to write while any of them is
missing.

    python analysis/apply_field_axis_not_validated.py --dry-run
    python analysis/apply_field_axis_not_validated.py --out-dir out_field

Run from the repository root.
"""
import argparse
import os

import docx

SRC = "out_blockers"
MS = "HT10016_revised_final2.docx"

OLD_101 = (
    "The applicability claim we carry forward is therefore graded rather than "
    "uniform. Iron chalcogenide 11-type is the only family that passes on the "
    "temperature axis, whose critical scale is directly reported, and it does "
    "not pass on the field axis. MgB2-class and iron pnictide 122-type pass on "
    "the field axis only, and their validation inherits the qualification of "
    "Section III.F. No family passes on both axes. Iron pnictide 1111-type "
    "fails on both axes and is excluded from candidate screening at "
    "substructure-aggregate scope. Table III records which axis validates each "
    "family, and the dispatch of Section III.E carries that label through to "
    "the predictions.")

NEW_101 = (
    "The applicability claim we carry forward is therefore graded rather than "
    "uniform, and on the field axis we no longer carry one. On the temperature "
    "axis, whose critical scale is directly reported, all three assessable "
    "families fall below the screening-grade threshold on the repaired cohort, "
    "at 0.546 for iron chalcogenide 11-type, 0.580 for iron pnictide 122-type "
    "and 0.513 for iron pnictide 1111-type, against 0.588, 1.314 and 3.120 on "
    "the cohort as published, where only the chalcogenide family passed. "
    "Section III.C states what part of that movement is a change of scale "
    "rather than of predictive skill, and the qualified claim made there is the "
    "one we carry. On the field axis we report no family-level verdict at all, "
    "for three reasons. The cohort that survives the anchor repairs and the "
    "stated fitting protocol is 52 fits from 12 papers rather than 94 from 16. "
    "The fitted field exponent does not separate substructure families on "
    "either cohort: the fraction of between-paper variance that the family "
    "label accounts for is 0.038 with a permutation p of 0.98 after repair and "
    "0.055 with p of 0.93 before it. And the per-family leave-one-compound-out "
    "errors reverse when the anchors are corrected, MgB2-class moving from "
    "0.753 to 1.230 and iron pnictide 122-type from 0.973 to 1.957 while iron "
    "chalcogenide 11-type moves from 1.093 to 0.707, so the two families that "
    "cleared the threshold stop clearing and the one that failed starts. A "
    "family-level verdict that changes sign when the scale behind it is "
    "corrected is not a verdict, and we withdraw the three we previously "
    "reported rather than replace them with their mirror image. Table III "
    "records the temperature-axis result for each family and records the field "
    "axis as not validated, and the dispatch of Section III.E carries that "
    "label through to the predictions.")

OLD_142 = (
    " They do carry the applicability validation for iron pnictide 122-type "
    "and MgB2-class, which is why Table IV labels the validating axis per "
    "family.")
NEW_142 = (
    " They no longer carry an applicability validation for any family: the "
    "field-axis verdicts this manuscript previously reported for iron pnictide "
    "122-type and MgB2-class are withdrawn in Section III.C, and Table IV "
    "labels the temperature axis alone.")

OLD_T3 = (
    "Field axis, on the cohort as published: MgB2-class 0.753 and 122-type "
    "0.973 clear; 11-type 1.093 and 1111-type 2.571 do not. On the repaired "
    "52-fit cohort those become 1.230, 1.957, 0.707 and 3.327, which reverses "
    "which families clear. That reversal is not yet carried through the "
    "family-level dispatch scope stated elsewhere in this paper, and the "
    "field-axis verdicts in this row remain those of the published cohort "
    "until it is.")
NEW_T3 = (
    "Field axis: not validated at family level, and no verdict is reported. "
    "The published values were MgB2-class 0.753 and 122-type 0.973 clearing "
    "with 11-type 1.093 and 1111-type 2.571 not clearing; on the repaired "
    "52-fit cohort they are 1.230, 1.957, 0.707 and 3.327, which reverses "
    "which families clear. The exponent also separates no family on either "
    "cohort, at 0.038 with p 0.98 after repair and 0.055 with p 0.93 before.")

# Passages the removal makes inconsistent. Located, printed, never edited: the
# dispatch scope is a claim about what the paper predicts, and changing it is
# not a correction.
FLAG = [
    "Dispatch is restricted to the three families that pass the validation and "
    "population requirements",
    "Only the MgB2 class dispatches under the refusal gates of this revision",
    "the reason the field-axis result is carried as a labelled limitation "
    "rather than as a validated prediction",
]


def replace_in_paragraph(p, find, repl):
    text = "".join(r.text for r in p.runs)
    if find not in text:
        return False
    start = text.index(find)
    end = start + len(find)
    pos, first, spans = 0, None, []
    for r in p.runs:
        a, b = pos, pos + len(r.text)
        if b > start and a < end:
            spans.append((r, max(start, a) - a, min(end, b) - a))
            if first is None:
                first = r
        pos = b
    if first is None:
        return False
    for r, i, j in spans:
        r.text = r.text[:i] + (repl if r is first else "") + r.text[j:]
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    d = docx.Document(os.path.join(SRC, MS))
    missed = []
    for find, repl, why in [(OLD_101, NEW_101, "the applicability claim"),
                            (OLD_142, NEW_142, "the Sec. III.F consequence")]:
        if not any(replace_in_paragraph(p, find, repl) for p in d.paragraphs):
            missed.append(why)
    cell = d.tables[2].rows[4].cells[2]
    if not any(replace_in_paragraph(p, OLD_T3, NEW_T3) for p in cell.paragraphs):
        missed.append("Table III, the field-axis clause")

    print("%s: 3 edits, %d not found" % (MS, len(missed)))
    for m in missed:
        print("   MISSING", m)

    print("\npassages the removal leaves inconsistent, located and not edited")
    for f in FLAG:
        hit = [i for i, p in enumerate(d.paragraphs) if f in p.text]
        print("   [%s] %s" % ("found" if hit else "NOT FOUND", f[:66]))
        if not hit:
            missed.append("flag: " + f[:40])
    print("   Section III.E restricts dispatch to families that pass "
          "validation, and the only")
    print("   family that dispatches is MgB2-class, whose validating axis was "
          "the field axis.")
    print("   With no field-axis verdict it has none. That is a decision about "
          "what the paper")
    print("   claims, not a correction, so it is left for the authors.")

    if missed:
        print("\nNothing written.")
        return 1
    if a.dry_run or not a.out_dir:
        print("\ndry run: nothing written")
        return 0
    os.makedirs(a.out_dir, exist_ok=True)
    out = os.path.join(a.out_dir, MS.replace("_final2", "_final3"))
    d.save(out)
    print("\nwritten: %s" % out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
