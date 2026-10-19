#!/usr/bin/env python3
"""
restore_acknowledgments.py

Put back the ACKNOWLEDGMENTS section, which analysis/apply_ai_disclosure.py
deleted.

What happened. That pass shortened the AI disclosure by replacing the first
paragraph of the section and removing every paragraph between the disclosure
heading and the heading it treated as the section's end, DATA AVAILABILITY.
In this manuscript ACKNOWLEDGMENTS sits between those two headings, so the
removal swallowed it: the heading and its paragraph, carrying two named
colleagues, an NSF award, a DARPA award and two compute allocations.

Why no check caught it. The run-time asserts compared the disclosure against
itself and the document text against its own new content. The review-item
sweep I had been using dropped the pin on "The authors acknowledge
illuminating discussions" when it was rewritten, so nothing was watching the
one string that would have failed.

The text is recovered verbatim from hossain/HT10016_final59_chowdhury.docx,
the version-59 manuscript, which predates every pass run today and is the
last copy of the section that exists. Nothing between version 59 and version
60 touched it.

The bounds bug is fixed in apply_ai_disclosure.py in the same commit: a
section now ends at the next heading of any kind, not at a named later one.
"""

import copy
import sys

import docx
from docx.oxml.ns import qn

SOURCE = "hossain/HT10016_final59_chowdhury.docx"
HEADING = "ACKNOWLEDGMENTS"
AFTER = "ARTIFICIAL INTELLIGENCE USE DISCLOSURE"
BEFORE = "DATA AVAILABILITY"

# asserted against the source, so a wrong recovery cannot pass quietly
MUST_CARRY = ["Alexander Balandin", "Aashman Mehta", "#2610404",
              "#D26AC50023", "PHY260151", "NAIRR260036"]


def recovered():
    doc = docx.Document(SOURCE)
    ps = doc.paragraphs
    i = next(k for k, p in enumerate(ps)
             if p.text.strip().upper().startswith(HEADING))
    body = next(p for p in ps[i + 1:i + 6] if p.text.strip())
    text = body.text.strip()
    for phrase in MUST_CARRY:
        assert phrase in text, "not in the recovered text: %s" % phrase
    return ps[i], body, text


def restore(path):
    doc = docx.Document(path)
    if any(p.text.strip().upper().startswith(HEADING) for p in doc.paragraphs):
        return False
    head_src, body_src, text = recovered()
    ps = doc.paragraphs
    target = next(p for p in ps if p.text.strip().upper().startswith(BEFORE))
    # style the restored heading like the one it sits beside, not like the
    # source document's, so it matches this document's own headings
    model_head = next(p for p in ps if p.text.strip().upper().startswith(AFTER))
    model_body = next(p for p in ps[ps.index(model_head) + 1:]
                      if p.text.strip())
    for src, model in ((head_src, model_head), (body_src, model_body)):
        new = copy.deepcopy(model._p)
        for node in list(new):
            if node.tag in (qn("w:r"), qn("w:hyperlink")):
                new.remove(node)
        run = copy.deepcopy(model._p.findall(qn("w:r"))[0])
        for t in run.findall(qn("w:t")):
            run.remove(t)
        for rPr in run.findall(qn("w:rPr")):
            for h in rPr.findall(qn("w:highlight")):
                rPr.remove(h)
        t = doc.element.makeelement(qn("w:t"), {})
        t.set(qn("xml:space"), "preserve")
        t.text = src.text.strip()
        run.append(t)
        new.append(run)
        target._p.addprevious(new)
        # the separating blank takes the BODY paragraph's style. Copying the
        # heading's instead leaves an empty Heading 1, which sets its own
        # spacing and prints as a gap the size of a heading.
        blank = copy.deepcopy(model_body._p)
        for node in list(blank):
            if node.tag in (qn("w:r"), qn("w:hyperlink")):
                blank.remove(node)
        target._p.addprevious(blank)
    doc.save(path)
    return True


def verify(path):
    doc = docx.Document(path)
    ps = [p.text.strip() for p in doc.paragraphs]
    assert HEADING in [t.upper() for t in ps], "no heading"
    i = next(k for k, t in enumerate(ps) if t.upper() == HEADING)
    body = next(t for t in ps[i + 1:i + 4] if t)
    for phrase in MUST_CARRY:
        assert phrase in body, phrase
    order = [k for k, t in enumerate(ps)
             if t.upper().startswith((AFTER, HEADING, BEFORE))]
    assert order == sorted(order) and len(order) == 3, order
    runs = 0
    for t in ps[i - 1:i + 3]:
        runs = runs + 1 if not t else 0
        assert runs < 2, "two blank paragraphs in a row"
    for k in (i - 1, i + 1):
        p = doc.paragraphs[k]
        if not p.text.strip():
            assert p.style.name == "Normal", (k, p.style.name)


if __name__ == "__main__":
    for path in ("out_final60/HT10016_revised_final60.docx",
                 "out_final60/HT10016_revised_final60_HIGHLIGHTED.docx"):
        did = restore(path)
        verify(path)
        print("%-46s acknowledgments %s"
              % (path.split("/")[-1], "restored" if did else "already present"))
