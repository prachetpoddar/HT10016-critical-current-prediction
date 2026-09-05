#!/usr/bin/env python3
"""Move the conditioning claim off the diagnostic that cannot carry it.

Two things the letter says against itself, both in one paragraph, and one gap
between the letter and the manuscript. All three are fixed here.

  Paragraph 38, answering the 23-fold, says the conditioning claim "no longer
  rests on this ratio at all. It rests on the variance-decomposition
  diagnostic". Paragraph 20, answering the MgB2 point, says that diagnostic
  "cannot establish the regime distinction it is used to draw". The claim is
  parked on the one piece of evidence the letter has already withdrawn,
  eighteen paragraphs earlier.

  The same paragraph says "all 96 per-paper anchor groups are identical before
  and after the magnetic-field unit repair". True of that repair, and
  superseded: the anchor repair withdrew 26 of the 96 and rescaled 4 more,
  which paragraph 20 states.

  And the field axis is reported as not validated at family level in Table III
  and Sec. III.C of the manuscript, and the letter never says so.

Where the conditioning claim actually rests, and the numbers written in:

  Temperature axis. The fraction of between-paper variance the substructure
  label accounts for in the fitted exponent, with papers as the unit and the
  family label shuffled across papers. On the papers the two cohorts share it
  goes from 0.436 to 0.524 with a permutation probability from 0.016 to 0.007.
  audit/headline_recomputed_20260905.md and audit/headline_on_repaired_cohort.csv.

  Field axis. 0.055 with a probability of 0.93 on the deposited passing cohort
  and 0.038 with 0.98 on the admitted one; matched on the twelve papers they
  share, 0.031 to 0.038 with 0.99 and 0.98. It separates nothing on either.

This is also the strengthening the temperature-axis result has been owed: the
letter has never mentioned it, and it is the claim both referees called
valuable.

Source out_a10/*_final9.docx, output out_basis/.

    python analysis/apply_conditioning_basis_edits.py --dry-run
    python analysis/apply_conditioning_basis_edits.py --out-dir out_basis

Run from the repository root.
"""
import argparse
import os

import docx

SRC = "out_a10"
RESP = "RESPONSE_TO_REFEREES_final9.docx"

OLD = (
    "The conditioning claim no longer rests on this ratio at all. It rests on "
    "the variance-decomposition diagnostic, which is computed on the "
    "critical-current anchor and uses neither a fitted exponent nor a critical "
    "field scale. We verified that quantity independently of the corrections: "
    "all 96 per-paper anchor groups are identical before and after the "
    "magnetic-field unit repair described below.")

NEW = (
    "The conditioning claim no longer rests on this ratio at all, and we have "
    "to be careful about what we move it onto. An earlier version of this "
    "reply rested it on the variance-decomposition diagnostic. That will not "
    "do, because we concede above that the diagnostic cannot establish the "
    "regime distinction it is used to draw: sample form is very nearly a "
    "relabelling of source paper in this corpus, and under the only null that "
    "respects that structure no family cell can reach a probability below "
    "0.10. We also said there that the 96 per-paper anchor groups behind it "
    "were unchanged by the magnetic-field unit repair, which was true of that "
    "repair and has since been overtaken: 26 of those 96 rows held values "
    "contradicting their source figures and are withdrawn, and 4 more carried "
    "a scale error and are corrected. "
    "What the claim rests on is the temperature axis, and we should have said "
    "so in our first reply rather than leaving the referee to find the "
    "strongest result unstated. Taking source papers as the unit and shuffling "
    "the substructure label between them, the fraction of between-paper "
    "variance in the fitted temperature exponent that the family label "
    "accounts for is 0.524, with a permutation probability of 0.007. On the "
    "papers shared by the cohorts before and after the transition-temperature "
    "repairs it rises from 0.436 to 0.524 and the probability falls from 0.016 "
    "to 0.007, so the repairs strengthen it rather than producing it. That is "
    "a family-level result on the axis whose critical scale is directly "
    "reported, and it is the one claim in this paper that improved under every "
    "correction we made. "
    "The field axis carries none of it, and we now say so plainly rather than "
    "reporting a weaker version. The same statistic on the field exponent is "
    "0.055 with a probability of 0.93 on the cohort as published and 0.038 "
    "with 0.98 on the 52 fits our stated protocol admits; matched on the "
    "twelve papers those two share, 0.031 against 0.038. The field exponent "
    "separates no substructure family on any cohort we can construct. Together "
    "with the per-family errors reversing when the anchors are repaired, that "
    "is why Table III now records the field axis as not validated at family "
    "level and reports no verdict for it.")


def replace_in_paragraph(p, find, repl):
    text = "".join(r.text for r in p.runs)
    if find not in text:
        return False
    start = text.index(find)
    end = start + len(find)
    pos, first, spans = 0, None, []
    for r in p.runs:
        a, b = pos, pos + len(r.text)
        if b > start and a < end:
            spans.append((r, max(start, a) - a, min(end, b) - a))
            if first is None:
                first = r
        pos = b
    if first is None:
        return False
    for r, i, j in spans:
        r.text = r.text[:i] + (repl if r is first else "") + r.text[j:]
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    d = docx.Document(os.path.join(SRC, RESP))
    if not any(replace_in_paragraph(p, OLD, NEW) for p in d.paragraphs):
        print("the paragraph was not found. Nothing written.")
        return 1
    print("%s: 1 edit applied" % RESP)

    full = " ".join(p.text for p in d.paragraphs)
    bad = []
    # The superseded verification must not survive unqualified.
    if "96 per-paper anchor groups are identical" in full and \
            "overtaken" not in full:
        bad.append("the superseded 96-anchor verification is unqualified")
    # The letter must now carry the field-axis position the manuscript takes.
    if "not validated at family level" not in full:
        bad.append("the letter still does not state the field-axis position")
    # And the temperature-axis result must actually be in it.
    if "0.524" not in full:
        bad.append("the temperature-axis result is still absent")
    if bad:
        print("\nNothing written.")
        for b in bad:
            print("  ", b)
        return 1
    print("the superseded verification is qualified, the field-axis position "
          "is stated,")
    print("and the temperature-axis result is in the letter for the first time")

    if a.dry_run or not a.out_dir:
        print("\ndry run: nothing written")
        return 0
    os.makedirs(a.out_dir, exist_ok=True)
    d.save(os.path.join(a.out_dir, RESP.replace("_final9", "_final10")))
    print("\nwritten to %s" % a.out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
