#!/usr/bin/env python3
"""
apply_data_availability.py

Cut the data-availability statement to two sentences, per Hossain.

It ran to five paragraphs, around 400 words, describing the deposit's
contents, its provenance labels, the withdrawn kilo-oersted fits, the four
scripts behind the dispatch table and the corpus enumeration. All of that is
now in the deposit's own README landing section, added on 2026-10-18, and in
Sec. 14 of the Supplemental Material, so the statement can do the one job a
data-availability statement has: say where the data is.

The first sentence names the deposit. The second says where the source
measurements came from, with the publisher and database sites cited, which is
the half a reader outside our institution needs in order to reach the
underlying articles.

Two things the old statement carried are not in the new one, and are listed
here rather than left to be noticed later:

  the named regeneration scripts for Table II, Sec. III.A and Sec. III.C,
  which the README now maps section by section;

  the commitment that "an archived snapshot with a permanent digital object
  identifier will be deposited before publication". That is a promise to the
  editor rather than a description of the data, and it needs a home. The
  cover letter is the natural place if it is not to sit here.

Every URL is linked, as the rest of the set already is.
"""

import copy
import sys

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, "analysis")
from apply_hyperlinks import isolate, link, urls_in  # noqa: E402

HEADING = "DATA AVAILABILITY"
NEXT = "REFERENCES"

NEW = (
    "The extracted critical-current dataset and the code that regenerates "
    "every number in this paper are publicly available at "
    "https://github.com/prachetpoddar/HT10016-critical-current-prediction. "
    "The source measurements are external to this work and are available "
    "from their publishers, Elsevier (https://www.sciencedirect.com), "
    "Springer Nature (https://link.springer.com), IOP Publishing "
    "(https://iopscience.iop.org) and arXiv (https://arxiv.org), together "
    "with datasets from the National High Magnetic Field Laboratory "
    "(https://nationalmaglab.org) and structural records from the Materials "
    "Project (https://materialsproject.org), each identified in our deposit."
)


def section_bounds(doc):
    """(index of the heading, index of the paragraph after the section)."""
    ps = doc.paragraphs
    start = next(i for i, p in enumerate(ps)
                 if p.text.strip().upper().startswith(HEADING))
    # the next heading of any kind, not the named one: stopping at a named
    # later heading deletes whatever sections sit between the two
    end = next(i for i, p in enumerate(ps[start + 1:], start + 1)
               if p.style.name.startswith("Heading") and p.text.strip())
    assert end > start
    return start, end


def apply(path, marked):
    doc = docx.Document(path)
    start, end = section_bounds(doc)
    body = doc.paragraphs[start + 1:end]
    first = next(p for p in body if p.text.strip())
    keep = first._p
    # one run carries the whole statement, taking the first run's formatting
    model = keep.findall(qn("w:r"))[0]
    for node in list(keep):
        if node.tag in (qn("w:r"), qn("w:hyperlink")):
            keep.remove(node)
    run = copy.deepcopy(model)
    for t in run.findall(qn("w:t")):
        run.remove(t)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = NEW
    run.append(t)
    keep.append(run)
    removed = 0
    for p in body:
        if p._p is keep:
            continue
        p._p.getparent().remove(p._p)
        removed += 1
    doc.save(path)

    # relink, and mark the new text in the editor's copy
    doc = docx.Document(path)
    start, end = section_bounds(doc)
    paragraph = next(p for p in doc.paragraphs[start + 1:end] if p.text.strip())
    text = "".join(r.text for r in paragraph.runs)
    links = 0
    for at, url in reversed(list(urls_in(text))):
        r = isolate(paragraph, at, len(url))
        assert r.text == url, (r.text, url)
        link(paragraph, r, url)
        links += 1
    if marked:
        for r in paragraph._p.iter(qn("w:r")):
            rPr = r.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                r.insert(0, rPr)
            for old in rPr.findall(qn("w:highlight")):
                rPr.remove(old)
            h = OxmlElement("w:highlight")
            h.set(qn("w:val"), "yellow")
            rPr.append(h)
    doc.save(path)
    return removed, links


def verify(path, marked):
    doc = docx.Document(path)
    start, end = section_bounds(doc)
    body = [p for p in doc.paragraphs[start + 1:end] if p.text.strip()]
    assert len(body) == 1, "%d paragraph(s) in the section" % len(body)
    text = body[0].text
    assert text == NEW, text[:80]
    anchors = ["".join(t.text or "" for t in h.iter(qn("w:t")))
               for h in body[0]._p.findall(qn("w:hyperlink"))]
    assert sorted(anchors) == sorted(u for _, u in urls_in(NEW)), anchors
    rels = [r.target_ref for r in doc.part.rels.values()
            if r.reltype.endswith("/hyperlink")]
    for a in anchors:
        assert a in rels, a
    if marked:
        yellow = "".join(
            t.text or ""
            for r in body[0]._p.iter(qn("w:r"))
            if r.find(qn("w:rPr")) is not None
            and r.find(qn("w:rPr")).find(qn("w:highlight")) is not None
            for t in r.iter(qn("w:t")))
        assert yellow == NEW, len(yellow)
    return len(anchors)


if __name__ == "__main__":
    for path, marked in (("out_final60/HT10016_revised_final60.docx", False),
                         ("out_final60/HT10016_revised_final60_HIGHLIGHTED.docx",
                          True)):
        removed, links = apply(path, marked)
        n = verify(path, marked)
        print("%-46s %d paragraph(s) removed, %d link(s), %d words"
              % (path.split("/")[-1], removed, n, len(NEW.split())))
