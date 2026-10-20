#!/usr/bin/env python3
"""
restore_inline_equation.py

Put the orphaned inline equation back in its sentence.

The paragraph-splitting pass divided one long paragraph into two by
character offset. An inline m:oMath is a direct child of w:p that
contributes no w:t text, so the split counted it as nothing and left it
sitting at the end of the first half, after the closing full stop, while the
sentence it belongs to moved to the second half.

The manuscript therefore reads

    all accepted MgB2-class records receive the same predicted value .

with the value gone, and the previous paragraph ends with a bare
log10Jc=4.980 after its own full stop. The submitted-with-v59 copy shows
where it belongs, between "predicted value " and the full stop.

This moves it there. It is a content repair, not a formatting one, so it
asserts both ends: the equation leaves the paragraph it was stranded in, and
the sentence reads as it does in v59.
"""

import copy
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MATH = "http://schemas.openxmlformats.org/officeDocument/2006/math"
OMATH = "{%s}oMath" % MATH

ANCHOR = "receive the same predicted value "
EXPECT = "log10Jc=4.980"


def run_text(r):
    return "".join(t.text or "" for t in r.iter(qn("w:t")))


def set_text(r, text):
    for t in r.findall(qn("w:t")):
        r.remove(t)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)


def flat(p_el):
    out = []
    for node in p_el:
        if node.tag == qn("w:r"):
            out.append(run_text(node))
        elif node.tag == OMATH:
            out.append("".join(t.text or "" for t in node.iter("{%s}t" % MATH)))
    return "".join(out)


def repair(path):
    doc = docx.Document(path)

    orphans = []
    for p in doc.paragraphs:
        text = flat(p._p)
        for node in p._p:
            if node.tag != OMATH:
                continue
            math = "".join(t.text or "" for t in node.iter("{%s}t" % MATH))
            # an inline equation whose paragraph is prose and which sits
            # after the paragraph's last full stop
            outside = "".join(t.text or "" for t in p._p.iter(qn("w:t")))
            if len(outside.split()) > 12 and text.rstrip().endswith(math):
                orphans.append((p, node, math))
    assert len(orphans) == 1, "expected one orphaned equation, found %d" % len(orphans)
    src, node, math = orphans[0]
    assert math == EXPECT, "unexpected equation %r" % math

    target = None
    for p in doc.paragraphs:
        if ANCHOR in flat(p._p):
            assert target is None, "the anchor sentence is not unique"
            target = p
    assert target is not None, "the anchor sentence is gone"

    # find the run holding the anchor and split it there
    at = flat(target._p).index(ANCHOR) + len(ANCHOR)
    pos = 0
    placed = False
    for child in list(target._p):
        if child.tag != qn("w:r"):
            continue
        t = run_text(child)
        if pos + len(t) < at:
            pos += len(t)
            continue
        cut = at - pos
        head, tail = t[:cut], t[cut:]
        others = [c for c in child if c.tag not in (qn("w:rPr"), qn("w:t"))]
        assert not others, "the anchor run carries %s" % others[0].tag
        if tail:
            after = copy.deepcopy(child)
            child.addnext(after)
            set_text(after, tail)
        set_text(child, head)
        src._p.remove(node)
        child.addnext(node)
        placed = True
        break
    assert placed, "no run held the anchor"

    doc.save(path)

    check = docx.Document(path)
    joined = [flat(p._p) for p in check.paragraphs]
    assert any(ANCHOR + EXPECT in t for t in joined), \
        "the equation did not land in the sentence"
    assert not any(t.rstrip().endswith(EXPECT) and len(t.split()) > 12
                   for t in joined), "the equation is still stranded"
    return True


if __name__ == "__main__":
    for p in sys.argv[1:]:
        repair(p)
        print("%s: inline equation restored to its sentence" % p)
