#!/usr/bin/env python3
"""Restore the subscripts my run-level edits flattened.

Every text replacement in this session puts the replacement into the first run
of the span it replaces and empties the rest, which keeps that run's formatting
and loses any variation inside the span. Across twelve passes that cost the
manuscript exactly three formulas: the chemical subscripts in SmFeAsO0.8F0.2
and in the Ba(Fe,Co)2As2 that replaced it, five subscript runs and nine
characters in Sec. III.F. The supplement and the response lost none, and no
italic, superscript or bold run was affected in any document.

The check that found it is not the obvious one. Counting runs carrying a
subscript flag finds nothing, because emptying a run's text leaves its font
attributes in place. Counting runs that carry the flag AND render is what
shows the loss.

There is a second, larger question this deliberately does not answer. The
manuscript already mixed the two spellings before this session: 15 flat
occurrences of MgB2 against 22 properly subscripted. That is the authors'
inconsistency and correcting it across the paper is not a correction, it is a
change of house style, so paragraphs this session never touched are left
exactly as they are.

What is fixed is only what this session did. Paragraphs are compared against
the pre-session document and only those whose text changed are rebuilt, which
covers the 3 subscripted formulas my edits flattened and the 8 flat ones my
replacement text introduced. In those paragraphs the formulas below are given
the dominant convention.

    python analysis/restore_formula_subscripts.py --dry-run
    python analysis/restore_formula_subscripts.py --out-dir out_send

Run from the repository root.
"""
import argparse
import copy
import os

import docx

SRC = "out_send"
MS = "HT10016_revised_final12.docx"

# formula as it appears in the flattened text, and the pieces that are
# subscript. Order matters: each piece is matched left to right.
FORMULAS = [
    ("Ba(Fe,Co)2As2", ["2", "2"]),
    ("SmFeAsO0.8F0.2", ["0.8", "0.2"]),
    ("MgB2", ["2"]),
]
BEFORE = os.path.join("out", "HT10016_revised_repaired.docx")


def split_formula(text, subs):
    """Text into (piece, is_subscript) parts, consuming subs left to right."""
    parts, i, rest = [], 0, list(subs)
    while i < len(text):
        if rest and text.startswith(rest[0], i):
            parts.append((rest[0], True))
            i += len(rest[0])
            rest.pop(0)
        else:
            if parts and not parts[-1][1]:
                parts[-1] = (parts[-1][0] + text[i], False)
            else:
                parts.append((text[i], False))
            i += 1
    return parts, not rest


def fix_paragraph(p, formula, subs):
    """Rebuild every occurrence of the formula in this paragraph, in place.

    Loops until no flat occurrence is left, because a single run can hold the
    formula several times and rebuilding splits the run.
    """
    done = 0
    while True:
        n = _fix_once(p, formula, subs)
        if not n:
            return done
        done += n


def _fix_once(p, formula, subs):
    done = 0
    for r in list(p.runs):
        if formula not in r.text or r.font.subscript:
            continue
        before, _s, after = r.text.partition(formula)
        parts, complete = split_formula(formula, subs)
        if not complete:
            continue
        r.text = before
        anchor = r
        for piece, is_sub in parts + ([(after, False)] if after else []):
            new = copy.deepcopy(r._r)
            anchor._r.addnext(new)
            nr = docx.text.run.Run(new, p)
            nr.text = piece
            nr.font.subscript = True if is_sub else None
            anchor = nr
        done += 1
        return done
    return done


def rendering_subscripts(d):
    return sum(len(r.text) for p in d.paragraphs for r in p.runs
               if r.text and r.font.subscript)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    d = docx.Document(os.path.join(SRC, MS))
    old = {p.text for p in docx.Document(BEFORE).paragraphs}
    touched = [p for p in d.paragraphs if p.text and p.text not in old]
    print("   paragraphs this session changed: %d of %d"
          % (len(touched), len(d.paragraphs)))
    before = rendering_subscripts(d)
    total = 0
    for formula, subs in FORMULAS:
        n = sum(fix_paragraph(p, formula, subs) for p in touched)
        print("   %-18s %d occurrence(s) rebuilt" % (formula, n))
        total += n
    after = rendering_subscripts(d)
    print("\n   subscript characters that render: %d before, %d after"
          % (before, after))

    ref = docx.Document(os.path.join("out", "HT10016_revised_repaired.docx"))
    target = rendering_subscripts(ref)
    print("   the version before this session's edits carried %d" % target)
    if total == 0:
        print("\n   nothing to do")
        return 0
    if after < before:
        print("\n   FAIL the rebuild removed subscripts")
        return 1
    if a.dry_run or not a.out_dir:
        print("\ndry run: nothing written")
        return 0
    d.save(os.path.join(a.out_dir, MS.replace("_final12", "_final13")))
    print("\n   written to %s" % a.out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
