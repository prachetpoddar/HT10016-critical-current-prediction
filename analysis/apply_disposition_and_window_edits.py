"""
apply_disposition_and_window_edits.py

Two disclosure repairs, plus the table the response letter promised.

**1. Eq. (1) reads "respectively" and Sec. III.C attributes 52 fits to "the
stated fitting protocol".** Applying Eq. (1) as printed returns 63: the
temperature condition is named for the temperature axis, and a reader of the
manuscript alone cannot get 52 from it. The response letter does disclose the
change, under its own heading, with the cost given to the fit ("11 above the
reduced temperature bound once the anchor is corrected"), but the manuscript
does not, and a referee reading the paper by itself finds an unreproducible
count. Sec. II.A now says which conditions are applied to which axis.

**2. The response letter says "the Supplemental Material now carries the
disposition of all 94 individually". It does not.** The supplement's sections
contain no such table, although audit/supplement_fit_disposition.csv holds all
94 rows. This module appends it as a new supplement section, so the promise and
the document agree.

Usage:
    python3 analysis/apply_disposition_and_window_edits.py IN_DIR OUT_DIR

The text edits are applied to the XML directly; the table is appended with
python-docx afterwards, because a 94-row table is not something to hand-write
into word/document.xml.
"""
import os
import shutil
import sys
import zipfile

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from apply_compound_label_edits import (paragraphs, plain,          # noqa: E402
                                        replace_in_paragraph, unescape)

MAIN = "HT10016_revised_final26.docx"
SUPP = "SUPPLEMENTAL_MATERIAL_revised_final26.docx"
RESP = "RESPONSE_TO_REFEREES_final26.docx"
OUT = {MAIN: "HT10016_revised_final27.docx",
       SUPP: "SUPPLEMENTAL_MATERIAL_revised_final27.docx",
       RESP: "RESPONSE_TO_REFEREES_final27.docx"}

DISPOSITION = os.path.join("audit", "supplement_fit_disposition.csv")
N_ROWS, N_KEPT, N_DROPPED = 94, 52, 42

EDITS = [
    (MAIN,
     "A partial fit is one in which the applicability window is satisfied on a "
     "single measurement axis but not both; it contributes that axis to the "
     "aggregation while the other axis remains uncharacterized for that "
     "record.",

     "A partial fit is one in which the applicability window is satisfied on a "
     "single measurement axis but not both; it contributes that axis to the "
     "aggregation while the other axis remains uncharacterized for that "
     "record. The word respectively above names which condition belongs to "
     "which axis, not which conditions are imposed on it. In this revision "
     "both conditions are imposed on both axes. In the original submission the "
     "temperature condition was imposed on the temperature axis only, and "
     "applying it to the field axis as well removes 11 of the 63 fits that "
     "clear the field condition by itself, which is how the 52-fit field "
     "cohort of Sec. III.C is reached. Supplement Sec. 16 gives the "
     "disposition of all 94 field-axis fits individually."),
]

SECTION_TITLE = "16. Disposition of the 94 field-axis fits"
SECTION_TEXT = (
    "Every field-axis Form 3 fit in the cohort as published, with the reason "
    "it is kept or dropped under the repaired anchors and the applicability "
    "window of Eq. (1) applied to both axes. Of the 94, %d are kept and %d are "
    "dropped: %d on the three papers whose critical-field anchors were "
    "withdrawn, %d for failing the field clause under the repaired anchor, and "
    "%d for a reduced temperature at or above 0.7 under the repaired "
    "transition temperature, which the original submission did not impose on "
    "this axis. The exponent shown is the value as deposited, before the "
    "repairs. Generated from audit/supplement_fit_disposition.csv."
)


def append_table(path):
    import docx
    d = pd.read_csv(DISPOSITION)
    if len(d) != N_ROWS:
        sys.exit("the disposition table holds %d rows, not the %d the "
                 "response letter promises" % (len(d), N_ROWS))
    counts = d.disposition.value_counts().to_dict()
    if counts.get("kept") != N_KEPT or counts.get("dropped") != N_DROPPED:
        sys.exit("the disposition table splits %s, not %d kept and %d dropped"
                 % (counts, N_KEPT, N_DROPPED))
    reasons = d[d.disposition == "dropped"].reason.astype(str)
    n_withdrawn = int(reasons.str.startswith("paper withdrawn").sum())
    n_field = int(reasons.str.contains("field clause").sum())
    n_temp = int(reasons.str.contains("T/Tc above").sum())
    if n_withdrawn + n_field + n_temp != N_DROPPED:
        sys.exit("the dropped reasons do not partition: %d + %d + %d against "
                 "%d" % (n_withdrawn, n_field, n_temp, N_DROPPED))

    doc = docx.Document(path)
    doc.add_paragraph()
    doc.add_paragraph(SECTION_TITLE)
    doc.add_paragraph(SECTION_TEXT % (N_KEPT, N_DROPPED, n_withdrawn, n_field,
                                      n_temp))
    head = ["Source", "Sample", "T (K)", "Deposited beta_H", "Disposition",
            "Reason"]
    t = doc.add_table(rows=1, cols=len(head))
    # The supplement defines exactly one table style, named "Table", which the
    # six tables already in it use. Asking for "Table Grid" raises, so the
    # style is taken from a table already in the document rather than named.
    if not doc.tables:
        sys.exit("the supplement carries no table to take a style from")
    t.style = doc.tables[0].style
    for c, h in enumerate(head):
        t.rows[0].cells[c].text = h
    for _, r in d.iterrows():
        cells = t.add_row().cells
        for c, v in enumerate((r.paper, r["sample"], "%.4g" % r.temperature_K,
                               "%.3f" % r.deposited_beta_H, r.disposition,
                               r.reason)):
            cells[c].text = str(v)
    doc.save(path)
    print("   appended %d rows to %s" % (len(d), os.path.basename(path)))


def main():
    if len(sys.argv) != 3:
        sys.exit("usage: apply_disposition_and_window_edits.py IN_DIR OUT_DIR")
    src, dst = sys.argv[1], sys.argv[2]
    if not os.path.exists(DISPOSITION):
        sys.exit("missing input: %s" % DISPOSITION)
    os.makedirs(dst, exist_ok=True)
    by_file = {}
    for f, find, repl in EDITS:
        by_file.setdefault(f, []).append((find, repl))
    for name in (MAIN, SUPP, RESP):
        ip, op = os.path.join(src, name), os.path.join(dst, OUT[name])
        if not os.path.exists(ip):
            sys.exit("missing input: %s" % ip)
        zin = zipfile.ZipFile(ip)
        doc = zin.read("word/document.xml").decode("utf-8")
        for find, repl in by_file.get(name, []):
            hits = [p for p in paragraphs(doc) if find in unescape(plain(p))]
            if len(hits) != 1:
                sys.exit("%s: %d paragraph(s) contain %r, expected exactly 1"
                         % (name, len(hits), find[:70]))
            new = replace_in_paragraph(hits[0], find, repl)
            if new is None:
                sys.exit("%s: could not place the replacement for %r"
                         % (name, find[:70]))
            doc = doc.replace(hits[0], new, 1)
            print("   %-8s %s" % (name.split("_")[0][:8], find[:62]))
        tmp = op + ".part"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/document.xml":
                    zout.writestr(item, doc)
                else:
                    zout.writestr(item, zin.read(item.filename))
        zin.close()
        shutil.move(tmp, op)
        print("   wrote %s" % op)
    append_table(os.path.join(dst, OUT[SUPP]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
